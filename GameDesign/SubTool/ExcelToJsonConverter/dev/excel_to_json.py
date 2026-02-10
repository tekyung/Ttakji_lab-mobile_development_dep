#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Excel to JSON Converter
엑셀 파일을 시트별 JSON 파일로 변환하는 GUI 툴
- 엑셀의 표(Table) 영역을 인식하여 해당 범위만 변환
"""

import json
import os
import re
import sys
import tkinter as tk
from tkinter import filedialog, scrolledtext, messagebox
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries


# Config 파일 설정
CONFIG_FILENAME = "config.json"


def get_script_dir() -> str:
    """스크립트 디렉토리 반환 (exe 빌드 시에도 동작)"""
    if getattr(sys, 'frozen', False):
        # PyInstaller로 빌드된 경우
        return os.path.dirname(sys.executable)
    else:
        # 일반 Python 실행
        return os.path.dirname(os.path.abspath(__file__))


def get_project_root() -> str:
    """프로젝트 루트 경로 반환"""
    script_dir = get_script_dir()
    # GameDesign/SubTool/ → 프로젝트 루트 (2단계 위)
    return os.path.dirname(os.path.dirname(script_dir))


def get_config_path() -> str:
    """Config 파일 경로 반환"""
    return os.path.join(get_script_dir(), CONFIG_FILENAME)


def load_config() -> Dict[str, str]:
    """Config 파일 로드 (상대경로 → 절대경로 변환)"""
    config_path = get_config_path()
    project_root = get_project_root()
    
    if not os.path.exists(config_path):
        return {}
    
    try:
        with open(config_path, 'r', encoding='utf-8') as f:
            config = json.load(f)
        
        # 상대경로 → 절대경로 변환
        result = {}
        for key, value in config.items():
            if value:
                # 상대경로인 경우 절대경로로 변환
                if not os.path.isabs(value):
                    abs_path = os.path.normpath(os.path.join(project_root, value))
                    result[key] = abs_path
                else:
                    result[key] = value
            else:
                result[key] = value
        
        return result
    except Exception:
        return {}


def save_config(input_path: str, output_path: str):
    """Config 파일 저장 (절대경로 → 상대경로 변환)"""
    config_path = get_config_path()
    project_root = get_project_root()
    
    config = {}
    
    # 절대경로 → 상대경로 변환
    if input_path:
        try:
            rel_path = os.path.relpath(input_path, project_root)
            config['input_path'] = rel_path.replace('\\', '/')
        except ValueError:
            # 다른 드라이브인 경우 절대경로 유지
            config['input_path'] = input_path
    
    if output_path:
        try:
            rel_path = os.path.relpath(output_path, project_root)
            config['output_path'] = rel_path.replace('\\', '/')
        except ValueError:
            config['output_path'] = output_path
    
    try:
        with open(config_path, 'w', encoding='utf-8') as f:
            json.dump(config, f, ensure_ascii=False, indent=2)
    except Exception:
        pass  # 저장 실패 시 무시


class ExcelToJsonConverter:
    """엑셀 파일을 JSON으로 변환하는 클래스"""
    
    # 정수형 타입 패턴
    INT_TYPES = {'int32', 'int16', 'byte'}
    
    # 문자열형 타입 패턴 (string(64), String(128) 등)
    STRING_PATTERN = re.compile(r'^[sS]tring\(\d+\)$')
    
    def __init__(self):
        self.log_callback = None
    
    def set_log_callback(self, callback):
        """로그 출력 콜백 설정"""
        self.log_callback = callback
    
    def log(self, message: str):
        """로그 메시지 출력"""
        if self.log_callback:
            self.log_callback(message)
        else:
            print(message)
    
    def is_int_type(self, type_str: str) -> bool:
        """정수형 타입인지 확인"""
        if type_str is None:
            return False
        return str(type_str).lower() in self.INT_TYPES
    
    def is_string_type(self, type_str: str) -> bool:
        """문자열형 타입인지 확인"""
        if type_str is None:
            return False
        return bool(self.STRING_PATTERN.match(str(type_str)))
    
    def convert_value(self, value: Any, type_str: str) -> Any:
        """자료형에 따라 값 변환"""
        # None/빈값 처리
        if value is None:
            return None
        
        # 정수형 변환
        if self.is_int_type(type_str):
            try:
                return int(float(value))
            except (ValueError, TypeError):
                return value
        
        # 문자열형 변환
        if self.is_string_type(type_str):
            return str(value)
        
        # Enum 타입 (E로 시작하는 타입) - 문자열로 유지
        if type_str and str(type_str).startswith('E'):
            if value is None or value == '':
                return None
            return str(value)
        
        # 기타: 숫자면 정수/실수, 아니면 문자열
        if isinstance(value, float):
            if value == int(value):
                return int(value)
            return value
        
        return value
    
    def is_comment_column(self, col_name: str) -> bool:
        """주석 컬럼인지 확인 (#으로 시작)"""
        if col_name is None:
            return True
        return str(col_name).startswith('#')
    
    def get_table_data(self, ws, table_ref: str) -> Tuple[List[str], List[str], List[List[Any]]]:
        """
        테이블 범위에서 데이터 추출
        Returns: (컬럼명 리스트, 자료형 리스트, 데이터 행 리스트)
        """
        min_col, min_row, max_col, max_row = range_boundaries(table_ref)
        
        # 모든 셀 데이터 읽기
        rows = []
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, 
                                 min_col=min_col, max_col=max_col, values_only=True):
            rows.append(list(row))
        
        if len(rows) < 2:
            return [], [], []
        
        # 첫 행: 컬럼명, 둘째 행: 자료형, 셋째 행부터: 데이터
        columns = rows[0]
        types = rows[1]
        data = rows[2:] if len(rows) > 2 else []
        
        return columns, types, data
    
    def process_table(self, ws, table_name: str, table_ref: str) -> Optional[List[Dict]]:
        """테이블 데이터를 JSON 형식으로 변환"""
        columns, types, data_rows = self.get_table_data(ws, table_ref)
        
        if not columns or not data_rows:
            self.log(f"    - {table_name}: 데이터 없음 (스킵)")
            return None
        
        # 유효한 컬럼 인덱스 필터링 (주석 컬럼 제외)
        valid_indices = []
        valid_columns = []
        valid_types = []
        
        for i, col in enumerate(columns):
            if not self.is_comment_column(col):
                valid_indices.append(i)
                valid_columns.append(col)
                valid_types.append(types[i] if i < len(types) else None)
        
        if not valid_columns:
            self.log(f"    - {table_name}: 유효한 컬럼 없음 (스킵)")
            return None
        
        # 데이터 변환
        records = []
        for row in data_rows:
            record = {}
            for idx, col_idx in enumerate(valid_indices):
                col_name = valid_columns[idx]
                type_str = valid_types[idx]
                value = row[col_idx] if col_idx < len(row) else None
                record[col_name] = self.convert_value(value, type_str)
            records.append(record)
        
        return records
    
    def convert_excel_file(self, excel_path: str, output_dir: str) -> List[str]:
        """엑셀 파일을 JSON으로 변환 (테이블 기반)"""
        created_files = []
        file_name = os.path.basename(excel_path)
        
        self.log(f"\n{file_name} 처리 중...")
        
        try:
            wb = load_workbook(excel_path, read_only=True, data_only=True)
        except Exception as e:
            self.log(f"  오류: 파일을 열 수 없습니다 - {e}")
            return created_files
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # 시트명을 테이블명으로 사용
            table_key = sheet_name
            
            # 시트 내 테이블 확인
            # read_only 모드에서는 tables 속성 접근 불가, 일반 모드로 다시 열기
            try:
                wb_full = load_workbook(excel_path, data_only=True)
                ws_full = wb_full[sheet_name]
                
                if not ws_full.tables:
                    self.log(f"    - {sheet_name}: 테이블 없음 (스킵)")
                    continue
                
                for table_name in ws_full.tables:
                    table_obj = ws_full.tables[table_name]
                    table_ref = table_obj.ref
                    
                    records = self.process_table(ws_full, table_name, table_ref)
                    
                    if records:
                        # JSON 파일 저장 (테이블 이름으로)
                        json_filename = f"{table_name}.json"
                        json_path = os.path.join(output_dir, json_filename)
                        
                        # 시트명을 키로 하는 객체 형식으로 저장
                        with open(json_path, 'w', encoding='utf-8') as f:
                            json.dump({table_key: records}, f, ensure_ascii=False, indent=2)
                        
                        created_files.append(json_filename)
                        self.log(f"    - {json_filename} 생성 ({len(records)}건)")
                
                wb_full.close()
                
            except Exception as e:
                self.log(f"    - {sheet_name}: 처리 오류 - {e}")
        
        wb.close()
        return created_files
    
    def convert_folder(self, input_dir: str, output_dir: str) -> Dict[str, List[str]]:
        """폴더 내 모든 엑셀 파일 변환"""
        results = {}
        
        # 출력 폴더 생성
        os.makedirs(output_dir, exist_ok=True)
        
        # 엑셀 파일 목록
        excel_files = [f for f in os.listdir(input_dir) 
                       if f.endswith(('.xlsx', '.xls')) and not f.startswith('~$')]
        
        if not excel_files:
            self.log("변환할 엑셀 파일이 없습니다.")
            return results
        
        self.log(f"총 {len(excel_files)}개의 엑셀 파일 발견")
        
        for excel_file in excel_files:
            excel_path = os.path.join(input_dir, excel_file)
            created_files = self.convert_excel_file(excel_path, output_dir)
            
            if created_files:
                results[excel_file] = created_files
        
        return results


class ConverterGUI:
    """변환 툴 GUI"""
    
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Excel to JSON Converter")
        self.root.geometry("600x500")
        self.root.resizable(True, True)
        
        self.converter = ExcelToJsonConverter()
        self.setup_ui()
        self.load_saved_config()
    
    def setup_ui(self):
        """UI 구성"""
        # 메인 프레임
        main_frame = tk.Frame(self.root, padx=10, pady=10)
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # 입력 폴더
        input_frame = tk.Frame(main_frame)
        input_frame.pack(fill=tk.X, pady=5)
        
        tk.Label(input_frame, text="입력 폴더:", width=10, anchor='w').pack(side=tk.LEFT)
        self.input_var = tk.StringVar()
        self.input_entry = tk.Entry(input_frame, textvariable=self.input_var)
        self.input_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        tk.Button(input_frame, text="찾아보기", command=self.browse_input).pack(side=tk.LEFT)
        
        # 출력 폴더
        output_frame = tk.Frame(main_frame)
        output_frame.pack(fill=tk.X, pady=5)
        
        tk.Label(output_frame, text="출력 폴더:", width=10, anchor='w').pack(side=tk.LEFT)
        self.output_var = tk.StringVar()
        self.output_entry = tk.Entry(output_frame, textvariable=self.output_var)
        self.output_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        tk.Button(output_frame, text="찾아보기", command=self.browse_output).pack(side=tk.LEFT)
        
        # 변환 버튼
        btn_frame = tk.Frame(main_frame)
        btn_frame.pack(fill=tk.X, pady=15)
        
        self.convert_btn = tk.Button(
            btn_frame, 
            text="변환 실행", 
            command=self.run_conversion,
            width=20,
            height=2
        )
        self.convert_btn.pack()
        
        # 로그 영역
        log_label = tk.Label(main_frame, text="로그 출력:", anchor='w')
        log_label.pack(fill=tk.X)
        
        self.log_text = scrolledtext.ScrolledText(
            main_frame, 
            height=15,
            state=tk.DISABLED
        )
        self.log_text.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # 컨버터에 로그 콜백 설정
        self.converter.set_log_callback(self.log)
    
    def load_saved_config(self):
        """저장된 config 로드 및 적용"""
        config = load_config()
        project_root = get_project_root()
        
        # config에서 경로 로드
        input_path = config.get('input_path', '')
        output_path = config.get('output_path', '')
        
        # config가 없거나 경로가 비어있으면 기본값 사용
        if not input_path:
            input_path = os.path.join(project_root, "Storage", "SampleData")
        if not output_path:
            output_path = os.path.join(project_root, "TCG_Project", "Data", "Generated")
        
        # 입력 경로가 존재하면 설정
        if os.path.exists(input_path):
            self.input_var.set(input_path)
        
        # 출력 경로의 부모 폴더가 존재하면 설정
        if os.path.exists(os.path.dirname(output_path)):
            self.output_var.set(output_path)
    
    def browse_input(self):
        """입력 폴더 선택"""
        folder = filedialog.askdirectory(title="엑셀 파일이 있는 폴더 선택")
        if folder:
            self.input_var.set(folder)
            # config 저장
            save_config(folder, self.output_var.get())
    
    def browse_output(self):
        """출력 폴더 선택"""
        folder = filedialog.askdirectory(title="JSON 파일 출력 폴더 선택")
        if folder:
            self.output_var.set(folder)
            # config 저장
            save_config(self.input_var.get(), folder)
    
    def log(self, message: str):
        """로그 메시지 출력"""
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)
        self.root.update()
    
    def clear_log(self):
        """로그 클리어"""
        self.log_text.config(state=tk.NORMAL)
        self.log_text.delete(1.0, tk.END)
        self.log_text.config(state=tk.DISABLED)
    
    def run_conversion(self):
        """변환 실행"""
        input_dir = self.input_var.get().strip()
        output_dir = self.output_var.get().strip()
        
        # 유효성 검사
        if not input_dir:
            messagebox.showerror("오류", "입력 폴더를 선택해주세요.")
            return
        
        if not output_dir:
            messagebox.showerror("오류", "출력 폴더를 선택해주세요.")
            return
        
        if not os.path.exists(input_dir):
            messagebox.showerror("오류", f"입력 폴더가 존재하지 않습니다:\n{input_dir}")
            return
        
        # 로그 클리어 및 변환 시작
        self.clear_log()
        self.log("=" * 50)
        self.log("Excel to JSON 변환 시작")
        self.log("=" * 50)
        self.log(f"입력: {input_dir}")
        self.log(f"출력: {output_dir}")
        
        # 버튼 비활성화
        self.convert_btn.config(state=tk.DISABLED)
        
        try:
            results = self.converter.convert_folder(input_dir, output_dir)
            
            # 결과 요약
            self.log("\n" + "=" * 50)
            self.log("변환 완료!")
            self.log("=" * 50)
            
            total_files = sum(len(files) for files in results.values())
            self.log(f"총 {len(results)}개 엑셀 파일에서 {total_files}개 JSON 파일 생성")
            
            if total_files > 0:
                messagebox.showinfo("완료", f"변환 완료!\n{total_files}개 JSON 파일 생성됨\n\n출력 폴더:\n{output_dir}")
            else:
                messagebox.showwarning("알림", "변환된 파일이 없습니다.\n엑셀 파일 구조를 확인해주세요.")
                
        except Exception as e:
            self.log(f"\n오류 발생: {e}")
            messagebox.showerror("오류", f"변환 중 오류가 발생했습니다:\n{e}")
        
        finally:
            # 버튼 활성화
            self.convert_btn.config(state=tk.NORMAL)
    
    def run(self):
        """GUI 실행"""
        self.root.mainloop()


def main():
    """메인 함수"""
    app = ConverterGUI()
    app.run()


if __name__ == "__main__":
    main()
